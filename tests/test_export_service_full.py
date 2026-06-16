"""Extended tests for ExportService — PDF, PPT, and edge cases."""
import csv
import io
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

from RJK.services.export_service import ExportService


@pytest.fixture
def svc():
    return ExportService()


SIMPLE = [
    {"NAME": "Alice", "SCORE": 95, "DEPT": "Sales"},
    {"NAME": "Bob",   "SCORE": 72, "DEPT": "Tech"},
]

SINGLE = [{"COL": "value"}]

DATE_ROWS = [
    {"PRODUCT": "Widget", "SALE_DATE": "2024-03-15", "REVENUE": 100.0},
    {"PRODUCT": "Gadget", "SALE_DATE": "2024-03-16", "REVENUE": 200.0},
]


# ---------------------------------------------------------------------------
# CSV edge cases (supplement existing tests)
# ---------------------------------------------------------------------------

class TestToCsvExtended:

    def test_single_row(self, svc):
        out = svc.to_csv(SINGLE).decode()
        lines = out.splitlines()
        assert len(lines) == 2
        assert "COL" in lines[0]

    def test_unicode_values(self, svc):
        rows = [{"CITY": "München"}, {"CITY": "São Paulo"}]
        out = svc.to_csv(rows).decode("utf-8")
        assert "München" in out
        assert "São Paulo" in out

    def test_special_chars_quoted(self, svc):
        rows = [{"NOTE": 'has,comma and "quotes"'}]
        out = svc.to_csv(rows).decode()
        reader = csv.DictReader(io.StringIO(out))
        assert list(reader)[0]["NOTE"] == 'has,comma and "quotes"'

    def test_none_value_rendered_empty(self, svc):
        rows = [{"A": None, "B": 1}]
        out = svc.to_csv(rows).decode()
        reader = csv.DictReader(io.StringIO(out))
        assert list(reader)[0]["A"] == ""


# ---------------------------------------------------------------------------
# Excel edge cases
# ---------------------------------------------------------------------------

class TestToExcelExtended:

    def test_date_string_coerced_to_date_cell(self, svc):
        import openpyxl
        from datetime import date
        out = svc.to_excel(DATE_ROWS)
        wb = openpyxl.load_workbook(io.BytesIO(out))
        ws = wb.active
        # Row 2 col 2 is SALE_DATE
        cell = ws.cell(2, 2)
        assert isinstance(cell.value, date)

    def test_numeric_values_preserved(self, svc):
        import openpyxl
        out = svc.to_excel(SIMPLE)
        wb = openpyxl.load_workbook(io.BytesIO(out))
        ws = wb.active
        assert ws.cell(2, 2).value == 95

    def test_sheet_name_truncated_to_31_chars(self, svc):
        import openpyxl
        long_name = "A" * 40
        out = svc.to_excel(SIMPLE, sheet_name=long_name)
        wb = openpyxl.load_workbook(io.BytesIO(out))
        assert len(wb.active.title) <= 31

    def test_single_column(self, svc):
        import openpyxl
        rows = [{"ONLY": i} for i in range(5)]
        out = svc.to_excel(rows)
        wb = openpyxl.load_workbook(io.BytesIO(out))
        assert wb.active.max_column == 1
        assert wb.active.max_row == 6  # header + 5 rows

    def test_many_rows(self, svc):
        import openpyxl
        rows = [{"IDX": i, "VAL": i * 2} for i in range(500)]
        out = svc.to_excel(rows)
        wb = openpyxl.load_workbook(io.BytesIO(out))
        assert wb.active.max_row == 501  # header + 500


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

class TestToPdf:

    def test_returns_bytes(self, svc):
        assert isinstance(svc.to_pdf(SIMPLE), bytes)

    def test_non_empty_output(self, svc):
        assert len(svc.to_pdf(SIMPLE)) > 0

    def test_starts_with_pdf_header(self, svc):
        assert svc.to_pdf(SIMPLE)[:4] == b"%PDF"

    def test_empty_rows_still_valid_pdf(self, svc):
        out = svc.to_pdf([])
        assert out[:4] == b"%PDF"

    def test_with_title(self, svc):
        out = svc.to_pdf(SIMPLE, title="My Report")
        assert len(out) > 0

    def test_with_params(self, svc):
        out = svc.to_pdf(SIMPLE, title="Report", params="date=2024-01-01")
        assert len(out) > 0

    def test_single_column(self, svc):
        rows = [{"COL": str(i)} for i in range(10)]
        out = svc.to_pdf(rows)
        assert out[:4] == b"%PDF"

    def test_many_columns(self, svc):
        rows = [{f"COL_{i}": i for i in range(20)}]
        out = svc.to_pdf(rows)
        assert out[:4] == b"%PDF"

    def test_large_dataset(self, svc):
        rows = [{"A": i, "B": i * 2, "C": f"row{i}"} for i in range(672)]
        out = svc.to_pdf(rows, title="Large Report")
        assert out[:4] == b"%PDF"
        assert len(out) > 10_000

    def test_special_chars_in_title(self, svc):
        out = svc.to_pdf(SIMPLE, title="Report & Summary <2024>")
        assert out[:4] == b"%PDF"

    def test_numeric_values_stringified(self, svc):
        rows = [{"AMOUNT": 1234567.89}]
        out = svc.to_pdf(rows)
        assert out[:4] == b"%PDF"


# ---------------------------------------------------------------------------
# PPT
# ---------------------------------------------------------------------------

class TestToPpt:

    def test_returns_bytes(self, svc):
        assert isinstance(svc.to_ppt(SIMPLE), bytes)

    def test_non_empty_output(self, svc):
        assert len(svc.to_ppt(SIMPLE)) > 0

    def test_valid_pptx_magic_bytes(self, svc):
        # PPTX is a ZIP — starts with PK
        assert svc.to_ppt(SIMPLE)[:2] == b"PK"

    def test_empty_rows_title_slide_only(self, svc):
        out = svc.to_ppt([])
        assert out[:2] == b"PK"

    def test_with_title_and_params(self, svc):
        out = svc.to_ppt(SIMPLE, title="Quarterly Report", params="region=North")
        assert out[:2] == b"PK"

    def test_readable_as_presentation(self, svc):
        from pptx import Presentation
        out = svc.to_ppt(SIMPLE, title="Test")
        prs = Presentation(io.BytesIO(out))
        assert prs.slides[0].shapes.title.text == "Test"

    def test_title_slide_present(self, svc):
        from pptx import Presentation
        out = svc.to_ppt(SIMPLE, title="My Title")
        prs = Presentation(io.BytesIO(out))
        assert len(prs.slides) >= 1

    def test_data_slides_created(self, svc):
        from pptx import Presentation
        rows = [{"A": i, "B": i} for i in range(10)]
        prs = Presentation(io.BytesIO(svc.to_ppt(rows)))
        # title slide + at least one data slide
        assert len(prs.slides) >= 2

    def test_large_dataset_multiple_slides(self, svc):
        from pptx import Presentation
        rows = [{"IDX": i, "VAL": i * 2} for i in range(100)]
        prs = Presentation(io.BytesIO(svc.to_ppt(rows)))
        # 100 rows / 35 per slide = 3 data slides + 1 title = 4
        assert len(prs.slides) >= 4

    def test_truncation_slide_added_over_500_rows(self, svc):
        from pptx import Presentation
        rows = [{"N": i} for i in range(600)]
        prs = Presentation(io.BytesIO(svc.to_ppt(rows)))
        titles = [s.shapes.title.text for s in prs.slides if s.shapes.title]
        assert any("truncated" in t.lower() for t in titles)

    def test_no_truncation_slide_under_500_rows(self, svc):
        from pptx import Presentation
        rows = [{"N": i} for i in range(100)]
        prs = Presentation(io.BytesIO(svc.to_ppt(rows)))
        titles = [s.shapes.title.text for s in prs.slides if s.shapes.title]
        assert not any("truncated" in t.lower() for t in titles)

    def test_large_dataset_672_rows(self, svc):
        rows = [{"F": f"F{i%8}", "R": f"R{i%6}", "V": i} for i in range(672)]
        out = svc.to_ppt(rows, title="Sales Report")
        assert out[:2] == b"PK"
        assert len(out) > 20_000

    def test_single_row(self, svc):
        from pptx import Presentation
        out = svc.to_ppt([{"COL": "only one row"}])
        prs = Presentation(io.BytesIO(out))
        assert len(prs.slides) >= 2
