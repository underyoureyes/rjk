import csv
import io

import pytest

from RJK.services.export_service import ExportService


@pytest.fixture
def service():
    return ExportService()


ROWS = [
    {"FLAVOUR": "Vanilla", "SCORE": 0.85, "COUNT": 1200},
    {"FLAVOUR": "Chocolate", "SCORE": 0.62, "COUNT": 800},
]


class TestToCsv:
    def test_returns_bytes(self, service):
        result = service.to_csv(ROWS)
        assert isinstance(result, bytes)

    def test_empty_rows(self, service):
        assert service.to_csv([]) == b""

    def test_header_present(self, service):
        lines = service.to_csv(ROWS).decode().splitlines()
        assert lines[0] == "FLAVOUR,SCORE,COUNT"

    def test_row_count(self, service):
        lines = service.to_csv(ROWS).decode().splitlines()
        # header + 2 data rows
        assert len(lines) == 3

    def test_values_correct(self, service):
        reader = csv.DictReader(io.StringIO(service.to_csv(ROWS).decode()))
        rows = list(reader)
        assert rows[0]["FLAVOUR"] == "Vanilla"
        assert float(rows[0]["SCORE"]) == pytest.approx(0.85)


class TestToExcel:
    def test_returns_bytes(self, service):
        result = service.to_excel(ROWS)
        assert isinstance(result, bytes)

    def test_empty_rows(self, service):
        result = service.to_excel([])
        assert isinstance(result, bytes)
        assert len(result) > 0  # valid empty workbook

    def test_correct_data(self, service):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(service.to_excel(ROWS)))
        ws = wb.active
        assert ws.cell(1, 1).value == "FLAVOUR"
        assert ws.cell(2, 1).value == "Vanilla"
        assert ws.cell(3, 1).value == "Chocolate"

    def test_sheet_name(self, service):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(service.to_excel(ROWS, sheet_name="MySheet")))
        assert wb.active.title == "MySheet"

    def test_header_fill_colour(self, service):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(service.to_excel(ROWS)))
        fill = wb.active.cell(1, 1).fill
        assert fill.fgColor.rgb.upper().endswith("003087")
