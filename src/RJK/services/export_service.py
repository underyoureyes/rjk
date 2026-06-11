import csv
import io
from datetime import date, datetime


def _coerce(val):
    """Convert ISO date strings to date objects so Excel writes real date cells."""
    if isinstance(val, str) and len(val) == 10:
        try:
            return datetime.strptime(val, "%Y-%m-%d").date()
        except ValueError:
            pass
    return val


class ExportService:
    def to_csv(self, rows: list[dict]) -> bytes:
        if not rows:
            return b""
        buf = io.StringIO()
        writer = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
        return buf.getvalue().encode("utf-8")

    def to_excel(self, rows: list[dict], sheet_name: str = "Report") -> bytes:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise RuntimeError("openpyxl is not installed") from exc

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name[:31]

        if not rows:
            buf = io.BytesIO()
            wb.save(buf)
            return buf.getvalue()

        headers = list(rows[0].keys())
        ws.append(headers)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="003087")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill

        for row in rows:
            ws.append([_coerce(row.get(h)) for h in headers])

        # Apply date format to columns that contain date objects
        for col_idx, header in enumerate(headers, start=1):
            col_vals = [row.get(header) for row in rows]
            if any(isinstance(_coerce(v), date) for v in col_vals):
                col_letter = get_column_letter(col_idx)
                for cell in ws[col_letter][1:]:   # skip header row
                    if isinstance(cell.value, date):
                        cell.number_format = "YYYY-MM-DD"

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
